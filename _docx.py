import docx, sys, os
p = r'd:\大四上\软件开发综合实验\小组提交文档-实验报告模板.docx'
d = docx.Document(p)
lines = []
for para in d.paragraphs:
    if para.text.strip():
        lines.append('[%s] %s' % (para.style.name, para.text.strip()))
for t_i, t in enumerate(d.tables):
    lines.append('--- TABLE %d ---' % t_i)
    for row in t.rows:
        cells = [c.text.strip().replace(chr(10), ' / ') for c in row.cells]
        lines.append(' | '.join(cells))
out = r'd:\大四上\软件开发综合实验\_txt\template.txt'
open(out, 'w', encoding='utf-8').write('\n'.join(lines))
print(len(lines))
